from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(r'C:\Users\Siddharth Gupta\Desktop\CODES-ALL\3dg')
PATH=ROOT/'output'/'PanoGS_Project_Report.docx'
doc=Document(PATH)

toc_items=[
('Chapter 1: Introduction and Background',1),('1.1 Background',2),('1.2 Project Overview',2),('1.3 Motivation',2),('1.4 Scope',2),('1.5 Report Organization',2),
('Chapter 2: Problem Definition and Objectives',1),('2.1 Problem Statement',2),('2.2 Aim',2),('2.3 Objectives',2),('2.4 Intended Users',2),('2.5 Assumptions',2),
('Chapter 3: Background and Core Concepts',1),('3.1 Equirectangular Projection and Spherical Rays',2),('3.2 Monocular Depth',2),('3.3 3D Gaussian Splatting',2),('3.4 LDI and Occlusion Layers',2),('3.5 Level of Detail',2),
('Chapter 4: Requirements Analysis',1),('4.1 Functional Requirements',2),('4.2 Non-Functional Requirements',2),('4.3 Hardware and Software Requirements',2),('4.4 Feasibility',2),
('Chapter 5: System Design and Architecture',1),('5.1 Architecture Overview',2),('5.2 Local Data Flow',2),('5.3 Separate Hosted Path',2),('5.4 Main Modules',2),('5.5 Project Screenshots: Architecture and Interface',2),
('Chapter 6: Reconstruction Methodology',1),('6.1 Panorama Input and Depth Estimation',2),('6.2 Aligning and Fusing Depth',2),('6.3 Room Priors and Depth Cleanup',2),('6.4 Back-Projection and Gaussian Initialization',2),('6.5 Filling Black or Empty Regions',2),('6.6 Video and Novel-View Prototypes',2),('6.7 Reconstruction Figures: Input, Depth and Output',2),
('Chapter 7: Implementation and Technology Stack',1),('7.1 AI Models and Weight Loading',2),('7.2 Training and Optimization',2),('7.3 API Keys and External Services',2),('7.4 Libraries and Their Uses',2),('7.5 Export and Viewing',2),
('Chapter 8: Testing and Evaluation',1),('8.1 Repository Tests',2),('8.2 What Current Metrics Measure',2),('8.3 How to Determine Which Method Is Best',2),
('Chapter 9: Results and Discussion',1),('9.1 Implemented Capabilities',2),('9.2 Documented Example Outputs',2),('9.3 Interpretation',2),('9.4 Visual Results and Benchmark Evidence',2),
('Chapter 10: Limitations, Future Scope and Conclusion',1),('10.1 Current Limitations',2),('10.2 Future Work',2),('10.3 Conclusion',2),
('References',1),('Appendix A: Running the Project',1),('A.1 Common Local Commands',2),('A.2 World Labs Key',2),('A.3 Operational Notes',2),('Appendix B: Repository Test Coverage',1)]

# Replace the short contents list with a compact, indented contents page.
paras=doc.paragraphs
toc=next(p for p in paras if p.text.strip()=='TABLE OF CONTENTS')
idx=paras.index(toc)
to_remove=[]
for p in paras[idx+1:]:
    if p.text.strip()=='LIST OF FIGURES AND TABLES': break
    if not p._p.xpath('.//w:br[@w:type="page"]'):
        to_remove.append(p)
for p in to_remove: p._element.getparent().remove(p._element)
toc.paragraph_format.space_after=Pt(2)
for label,level in toc_items:
    p=doc.add_paragraph()
    p.paragraph_format.left_indent=Inches(.0 if level==1 else .22)
    p.paragraph_format.first_line_indent=Inches(0)
    p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(1)
    p.paragraph_format.line_spacing=1.0
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.35),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
    r=p.add_run(('• ' if level==1 else '   o ')+label+'\tXX')
    r.font.name='Times New Roman'; r.font.size=Pt(8.7 if level==2 else 9.1); r.bold=(level==1)
# Place the new TOC rows between the TOC title and the page break before the figures list.
toc_rows=[p for p in doc.paragraphs if p.text.endswith('\tXX')]
figures=next(p for p in doc.paragraphs if p.text.strip()=='LIST OF FIGURES AND TABLES')
figures_break=figures._p.getprevious()
for p in toc_rows:
    p._p.getparent().remove(p._p)
for p in toc_rows:
    figures_break.addprevious(p._p)

def add_slot_before(chapter_no,title,description,caption):
    heading=next(p for p in doc.paragraphs if p.style.name=='Heading 1' and p.text.startswith(f'CHAPTER {chapter_no}\n'))
    anchor=heading._p.getprevious() # chapter's existing page-break paragraph
    if anchor is None: raise RuntimeError('chapter page break not found')
    elements=[]
    pb=doc.add_paragraph(); pb.add_run().add_break(WD_BREAK.PAGE); elements.append(pb._p)
    hp=doc.add_paragraph(style='Heading 2'); hp.alignment=WD_ALIGN_PARAGRAPH.CENTER; hp.paragraph_format.keep_with_next=True; hp.add_run(title); elements.append(hp._p)
    cp=doc.add_paragraph(description); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.first_line_indent=Inches(0); cp.paragraph_format.space_after=Pt(10); elements.append(cp._p)
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    cell=t.cell(0,0); cell.width=Inches(6.35); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.text='\n\n\n\n\n\n\n\n\n\nPASTE PROJECT IMAGE HERE\n\n\n\n\n\n\n\n\n\n'
    for p in cell.paragraphs:
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name='Times New Roman'; r.font.size=Pt(10); r.font.color.rgb=RGBColor(120,120,120)
    trPr=t.rows[0]._tr.get_or_add_trPr(); cant=OxmlElement('w:cantSplit'); trPr.append(cant)
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'dashed'); e.set(qn('w:sz'),'12'); e.set(qn('w:color'),'7B8DA3'); borders.append(e)
    t._tbl.tblPr.append(borders)
    elements.append(t._tbl)
    cp2=doc.add_paragraph(caption); cp2.alignment=WD_ALIGN_PARAGRAPH.CENTER; cp2.paragraph_format.first_line_indent=Inches(0); cp2.paragraph_format.space_before=Pt(8); cp2.paragraph_format.space_after=Pt(0); cp2.runs[0].italic=True; elements.append(cp2._p)
    for el in elements: anchor.addprevious(el)

add_slot_before(6,'5.5 Project Screenshots: Architecture and Interface','Insert a screenshot of the PanoGS interface, architecture, or CLI workflow. Replace this frame with the project image before submission.','Figure 5.2  Project interface or architecture screenshot')
add_slot_before(7,'6.7 Reconstruction Figures: Input, Depth and Output','Suggested images: source panorama, predicted depth/heatmap, and reconstructed scene. Use this page for a labeled comparison.','Figure 6.2  Input panorama, depth estimate, and reconstruction')
add_slot_before(10,'9.4 Visual Results and Benchmark Evidence','Insert rendered novel views, scene screenshots, or benchmark evidence. State the scene, settings, and metric beside each result.','Figure 9.1  PanoGS scene result and evaluation evidence')

doc.save(PATH)
print(PATH)
