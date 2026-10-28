from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(r'C:\Users\Siddharth Gupta\Desktop\CODES-ALL\3dg')
OUT=ROOT/'output'; TMP=ROOT/'tmp'; OUT.mkdir(exist_ok=True)
DOCX=OUT/'PanoGS_Project_Report.docx'
NAVY='203957'; BLUE='E6EDF5'; PALE='F5F7FA'; EDGE='BAC5D2'

def border_table(t):
    pr=t._tbl.tblPr; b=pr.first_child_found_in('w:tblBorders')
    if b is None: b=OxmlElement('w:tblBorders'); pr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=b.find(qn('w:'+edge))
        if e is None: e=OxmlElement('w:'+edge); b.append(e)
        e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'5'); e.set(qn('w:color'),EDGE)
def shade(c,color):
    pr=c._tc.get_or_add_tcPr(); s=OxmlElement('w:shd'); s.set(qn('w:fill'),color); pr.append(s)
def table(doc,heads,rows,widths=None,size=9):
    t=doc.add_table(rows=1,cols=len(heads)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for j,x in enumerate(heads): t.rows[0].cells[j].text=str(x)
    for row in rows:
        c=t.add_row().cells
        for j,x in enumerate(row): c[j].text=str(x)
    border_table(t)
    for ri,row in enumerate(t.rows):
        if ri==0:
            tr=OxmlElement('w:tblHeader'); tr.set(qn('w:val'),'true'); row._tr.get_or_add_trPr().append(tr)
        for ci,c in enumerate(row.cells):
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths: c.width=Inches(widths[ci])
            shade(c,NAVY if ri==0 else (PALE if ri%2==0 else 'FFFFFF'))
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(2); p.paragraph_format.line_spacing=1.03
                for r in p.runs:
                    r.font.name='Times New Roman'; r.font.size=Pt(size)
                    if ri==0: r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t
def para(doc,text='',style=None,align=None,bold_prefix=None):
    p=doc.add_paragraph(style=style)
    if align is not None: p.alignment=align
    if bold_prefix and text.startswith(bold_prefix): p.add_run(bold_prefix).bold=True; p.add_run(text[len(bold_prefix):])
    else: p.add_run(text)
    return p
def h(doc,text,level=2):
    p=doc.add_paragraph(text,style=f'Heading {level}'); p.paragraph_format.keep_with_next=True; return p
def bullet(doc,text):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.left_indent=Inches(.28); p.paragraph_format.first_line_indent=Inches(-.15); p.paragraph_format.space_after=Pt(3); p.add_run(text); return p
def numbered(doc,text):
    p=doc.add_paragraph(style='List Number'); p.paragraph_format.left_indent=Inches(.3); p.paragraph_format.first_line_indent=Inches(-.16); p.paragraph_format.space_after=Pt(4); p.add_run(text); return p
def chapter(doc,n,title):
    doc.add_page_break(); p=doc.add_paragraph(style='Heading 1'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next=True; p.add_run(f'CHAPTER {n}\n{title.upper()}')
def page_field(p):
    r=p.add_run(); a=OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'),'begin'); i=OxmlElement('w:instrText'); i.set(qn('xml:space'),'preserve'); i.text=' PAGE '; z=OxmlElement('w:fldChar'); z.set(qn('w:fldCharType'),'end'); r._r.extend([a,i,z])
def diagram(path,title,labels):
    W,H=1500,560; im=Image.new('RGB',(W,H),'white'); d=ImageDraw.Draw(im)
    f=r'C:\Windows\Fonts\arial.ttf'; fb=r'C:\Windows\Fonts\arialbd.ttf'
    normal=ImageFont.truetype(f,19); bold=ImageFont.truetype(fb,23)
    d.text((W//2,38),title,anchor='mm',fill='#203957',font=bold)
    margin,gap=42,22; bw=(W-2*margin-gap*(len(labels)-1))//len(labels); y,bh=145,250
    for i,(name,detail) in enumerate(labels):
        x=margin+i*(bw+gap); d.rounded_rectangle((x,y,x+bw,y+bh),radius=20,fill='#E6EDF5',outline='#35577D',width=3)
        d.multiline_text((x+bw/2,y+53),name,anchor='mm',align='center',fill='#203957',font=bold,spacing=6)
        words=detail.split(); lines=[]; line=''
        for word in words:
            test=(line+' '+word).strip()
            if d.textbbox((0,0),test,font=normal)[2]>bw-34: lines.append(line); line=word
            else: line=test
        if line: lines.append(line)
        d.multiline_text((x+bw/2,y+150),'\n'.join(lines[:4]),anchor='mm',align='center',fill='#263343',font=normal,spacing=5)
        if i<len(labels)-1:
            ex=x+bw+gap-5; cy=y+bh//2; d.line((x+bw+4,cy,ex,cy),fill='#536A84',width=4); d.polygon([(ex,cy),(ex-13,cy-9),(ex-13,cy+9)],fill='#536A84')
    d.text((W//2,H-32),'Source-derived architecture schematic; not a runtime screenshot.',anchor='mm',fill='#687386',font=normal)
    im.save(path)

flow=TMP/'panogs_flow.png'; diagram(flow,'Local Panorama Reconstruction',[('Image input','Equirectangular panorama, RGB or HDR radiance'),('Depth inference','Six overlapping perspective faces; pretrained DA V2 or MiDaS'),('Geometry','Face alignment, spherical rays, room prior and optional LDI'),('Gaussian scene','Back-projected points initialized as anisotropic splats'),('Export and view','PLY, SPLAT, HDR SPLAT; PanoGS viewer or SuperSplat')])
arch=TMP/'panogs_architecture.png'; diagram(arch,'PanoGS Components',[('Interfaces','CLI and local Studio upload'),('Core services','Image and video I/O, depth models, camera geometry'),('Reconstruction','Layout, point cloud, filtering, inpainting, LDI'),('Scene tools','Gaussian initialization, optimization, compression, export'),('Viewers','Three.js WebGL viewer; separate SuperSplat editor')])

doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.78); sec.bottom_margin=Inches(.75); sec.left_margin=Inches(.9); sec.right_margin=Inches(.9); sec.header_distance=Inches(.35); sec.footer_distance=Inches(.35); sec.different_first_page_header_footer=True
normal=doc.styles['Normal']; normal.font.name='Times New Roman'; normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman'); normal.font.size=Pt(11); normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; normal.paragraph_format.line_spacing=1.18; normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.first_line_indent=Inches(.22)
for nm,sz in [('Title',22),('Heading 1',16),('Heading 2',12.5),('Heading 3',11)]:
    s=doc.styles[nm]; s.font.name='Times New Roman'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman'); s.font.size=Pt(sz); s.font.bold=True; s.font.color.rgb=RGBColor(0,0,0); s.paragraph_format.space_before=Pt(10); s.paragraph_format.space_after=Pt(7); s.paragraph_format.keep_with_next=True
if 'Small Note' not in doc.styles:
    s=doc.styles.add_style('Small Note',WD_STYLE_TYPE.PARAGRAPH); s.font.name='Times New Roman'; s.font.size=Pt(9); s.font.italic=True; s.font.color.rgb=RGBColor(80,80,80); s.paragraph_format.first_line_indent=Inches(0); s.paragraph_format.space_after=Pt(5)
s=doc.styles['Caption']; s.font.name='Times New Roman'; s.font.size=Pt(9); s.font.italic=True; s.font.color.rgb=RGBColor(45,45,45)
hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; rr=hp.add_run('PanoGS  |  Project Report'); rr.font.name='Times New Roman'; rr.font.size=Pt(9); rr.font.color.rgb=RGBColor(90,90,90)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; page_field(fp)

# Cover
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(58); p.paragraph_format.space_after=Pt(9); r=p.add_run('A PROJECT REPORT'); r.bold=True; r.font.size=Pt(17)
para(doc,'on',align=WD_ALIGN_PARAGRAPH.CENTER).runs[0].italic=True
p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(20); p.add_run('PanoGS An Engine for Converting Single Panoramic Images into 3D Gaussian Scenes')
p=para(doc,'Submitted by',align=WD_ALIGN_PARAGRAPH.CENTER); p.runs[0].bold=True; p.runs[0].italic=True
for name,reg in [('ABEL BOBBY','RA2411003030217'),('SIDDHARTH GUPTA','RA2411003030223'),('ADITYA SHARMA','RA2411003030234')]:
    p=para(doc,f'{name}  [Reg No: {reg}]',align=WD_ALIGN_PARAGRAPH.CENTER); p.paragraph_format.space_after=Pt(2); p.runs[0].bold=True
p=para(doc,'Under the guidance of',align=WD_ALIGN_PARAGRAPH.CENTER); p.runs[0].italic=True
p=para(doc,'Dr. Nishant Kumar Singh',align=WD_ALIGN_PARAGRAPH.CENTER); p.runs[0].bold=True
para(doc,'(Assistant Professor, Department of Computer Science & Engineering)',align=WD_ALIGN_PARAGRAPH.CENTER)
for text in ['in partial fulfillment for the award of the degree of','BACHELOR OF TECHNOLOGY','COMPUTER SCIENCE & ENGINEERING','FACULTY OF ENGINEERING AND TECHNOLOGY']:
    p=para(doc,text,align=WD_ALIGN_PARAGRAPH.CENTER); p.paragraph_format.space_after=Pt(3)
    if text.isupper(): p.runs[0].bold=True
# Extract and reuse the institution logo from the user-supplied reference PDF.
logo=TMP/'srm_logo.png'
try:
    from pypdf import PdfReader
    imgs=PdfReader(r'C:\Users\Siddharth Gupta\Downloads\bonafide-correctn (1).pdf').pages[0].images
    if imgs:
        logo.write_bytes(imgs[0].data); li=Image.open(logo).convert('RGBA'); li.thumbnail((450,150)); li.save(logo)
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(logo),width=Inches(2.6))
except Exception: pass
para(doc,'SRM INSTITUTE OF SCIENCE & TECHNOLOGY, DELHI NCR CAMPUS',align=WD_ALIGN_PARAGRAPH.CENTER)
p=para(doc,'October 2026',align=WD_ALIGN_PARAGRAPH.CENTER); p.runs[0].bold=True
p=para(doc,'SRM INSTITUTE OF SCIENCE & TECHNOLOGY',align=WD_ALIGN_PARAGRAPH.CENTER); p.runs[0].bold=True
para(doc,'(Under Section 3 of UGC Act, 1956)',align=WD_ALIGN_PARAGRAPH.CENTER)

# Certificate draft, with no fabricated signatures.
doc.add_page_break(); p=para(doc,'BONAFIDE CERTIFICATE',align=WD_ALIGN_PARAGRAPH.CENTER); p.runs[0].bold=True; p.runs[0].font.size=Pt(16)
para(doc,'This is to certify that the project report titled “PanoGS: An Engine for Converting Single Panoramic Images into 3D Gaussian Scenes” is submitted by Abel Bobby (RA2411003030217), Siddharth Gupta (RA2411003030223), and Aditya Sharma (RA2411003030234) in partial fulfillment of the requirements for the Bachelor of Technology degree in Computer Science & Engineering at SRM Institute of Science & Technology, Delhi NCR Campus.')
para(doc,'This certificate page is prepared for review and institutional completion. The signatures and approval below must be supplied by the authorized faculty.')
doc.add_paragraph(); doc.add_paragraph()
t=doc.add_table(rows=1,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
for c,txt in zip(t.rows[0].cells,['____________________________\nDr. Nishant Kumar Singh\nProject Guide','____________________________\nDr. Neelam Singh\nHead of Department']): c.text=txt; c.width=Inches(3); c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
border_table(t)

# Acknowledgement and abstract.
doc.add_page_break(); p=para(doc,'ACKNOWLEDGEMENT',align=WD_ALIGN_PARAGRAPH.CENTER); p.runs[0].bold=True
para(doc,'We express our gratitude to our project guide, Dr. Nishant Kumar Singh, for his guidance and feedback during the development of this project. We also thank the Head of the Department, Dr. Neelam Singh, and the faculty members of the Department of Computer Science & Engineering, SRM Institute of Science & Technology, Delhi NCR Campus, for the academic support and facilities provided.')
para(doc,'We are grateful to our families and friends for their encouragement throughout the project.')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; p.paragraph_format.space_before=Pt(18); p.add_run('Abel Bobby\nSiddharth Gupta\nAditya Sharma')
doc.add_page_break(); p=para(doc,'ABSTRACT',align=WD_ALIGN_PARAGRAPH.CENTER); p.runs[0].bold=True
para(doc,'PanoGS is a Python-based system for converting panoramas and related visual inputs into interactive 3D Gaussian scenes. Its local panorama workflow estimates monocular depth from overlapping perspective views, combines depth with spherical camera geometry and optional indoor-room priors, back-projects image samples into a point cloud, initializes Gaussian primitives, and exports PLY and splat formats. The project also includes image inspection, a video reconstruction prototype, point-cloud filtering, scene optimization, compression, a custom browser viewer, and a separate bundled SuperSplat editor.')
para(doc,'The depth estimators use pretrained Depth Anything V2 Metric Indoor or MiDaS weights. PanoGS performs inference on each submitted input; it does not train those depth networks. Its optional training command instead optimizes the Gaussian parameters of a particular scene against perspective crops sampled from the source image. An independent World Labs Marble workflow can upload an image preview and download a hosted-generated Gaussian scene.')
para(doc,'This report explains the architecture, algorithms, model loading, API-key handling, inpainting, LDI, LOD, and evaluation code. It distinguishes implemented processing from inferred geometry and from future evaluation. The repository contains unit-test source files and documents example outputs, but no accuracy result here is presented as a verified benchmark. The benchmark implementation compares adjacent rendered frames for PSNR and SSIM rather than comparing them with ground-truth views. A controlled held-out evaluation is required before claiming the system or a model is best.')
para(doc,'Keywords: panoramic reconstruction, monocular depth, 3D Gaussian Splatting, point cloud, inpainting, level of detail, PanoGS.')

# Contents and list of illustrations.
doc.add_page_break(); p=para(doc,'TABLE OF CONTENTS',align=WD_ALIGN_PARAGRAPH.CENTER); p.runs[0].bold=True
for x in ['Chapter 1  Introduction and Background','Chapter 2  Problem Definition and Objectives','Chapter 3  Background and Core Concepts','Chapter 4  Requirements Analysis','Chapter 5  System Design and Architecture','Chapter 6  Reconstruction Methodology','Chapter 7  Implementation and Technology Stack','Chapter 8  Testing and Evaluation','Chapter 9  Results and Discussion','Chapter 10  Limitations, Future Scope and Conclusion','References','Appendix A  Running the Project','Appendix B  Repository Test Coverage']:
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.25); p.paragraph_format.first_line_indent=Inches(0); p.paragraph_format.space_after=Pt(7); p.add_run(x)
p=doc.add_paragraph(style='Small Note'); p.add_run('Page numbers may change as this report is edited.')
doc.add_page_break(); p=para(doc,'LIST OF FIGURES AND TABLES',align=WD_ALIGN_PARAGRAPH.CENTER); p.runs[0].bold=True
for x in ['Figure 5.1  PanoGS components and data flow','Figure 6.1  Local panorama reconstruction flow','Table 3.1  Depth model roles','Table 4.1  Functional requirements','Table 4.2  Software and hardware requirements','Table 6.1  Reconstruction stages','Table 7.1  Libraries and their uses','Table 8.1  Evaluation measures and current status','Table B.1  Repository test modules']:
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.22); p.paragraph_format.first_line_indent=Inches(0); p.paragraph_format.space_after=Pt(5); p.add_run(x)

# Chapter 1
chapter(doc,1,'Introduction and Background'); h(doc,'1.1 Background')
para(doc,'A panoramic image captures a wide field of view, but stores appearance as a two-dimensional equirectangular raster. It does not directly contain a 3D mesh, camera motion, or observations of surfaces hidden from the capture point. PanoGS addresses this gap by estimating depth and assigning each image sample a 3D position, then representing the resulting scene with Gaussian primitives that can be rendered from nearby viewpoints.')
h(doc,'1.2 Project Overview'); para(doc,'PanoGS is a local-first Python package and command-line application. Its main path is panorama-to-depth-to-point-cloud-to-Gaussians. The package also provides commands for inspecting images, exporting depth visualizations, processing point clouds, reconstructing from video, optimizing a Gaussian PLY, compressing a scene, rendering a CPU preview, benchmarking, and viewing a scene.')
h(doc,'1.3 Motivation'); para(doc,'A single panorama is an accessible input for an immersive look-around experience. The engineering challenge is to turn spherical image coordinates and uncertain monocular depth into a continuous scene while managing seams, outliers, occlusion holes, large splat counts, and browser rendering cost.')
h(doc,'1.4 Scope'); para(doc,'This report covers repository source, documented commands and model integrations, and test files present in the repository. It does not claim that one panorama yields complete hidden geometry. The optional World Labs workflow is separate from local reconstruction. Notebook experiments are identified as external baselines rather than core CLI capabilities.')
h(doc,'1.5 Report Organization'); para(doc,'Chapters 2–4 define the problem, concepts, and requirements. Chapters 5–7 describe architecture, reconstruction methods, models, libraries, and training. Chapters 8–9 discuss tests and evaluation evidence. Chapter 10 summarizes limitations and next steps.')

# Chapter 2
chapter(doc,2,'Problem Definition and Objectives'); h(doc,'2.1 Problem Statement')
para(doc,'A single panorama gives color in every direction but usually lacks direct range measurements and parallax observations. Converting it to a navigable scene requires estimates of surface distance, consistent projection across the sphere, an efficient scene representation, and a renderer. Geometry inferred behind occluders or away from the camera center is uncertain.')
h(doc,'2.2 Aim'); para(doc,'To design and implement a local pipeline that transforms a panoramic image into a viewable 3D Gaussian scene using pretrained depth inference and explicit geometric processing.')
h(doc,'2.3 Objectives')
for x in ['Load common image formats and preserve HDR radiance where supported.','Provide selectable monocular-depth estimators behind a shared interface.','Reduce panorama distortion by estimating depth on overlapping perspective cubemap faces.','Fuse depth with spherical rays and indoor layout priors to form 3D points.','Handle depth edges and selected occlusion regions with classical image and geometry processing.','Initialize and export Gaussian scenes in PLY, SPLAT, and HDR SPLAT forms.','Provide local rendering, viewing, compression, and per-scene optimization tools.','Document limitations and establish a repeatable evaluation plan.']: bullet(doc,x)
h(doc,'2.4 Intended Users'); para(doc,'The system is intended for students, developers, and researchers exploring panorama reconstruction, depth estimation, Gaussian scene formats, and local visualization. The World Labs command is for users who explicitly choose hosted generation and have a provider key and available credits.')
h(doc,'2.5 Assumptions'); para(doc,'The panorama is assumed to be equirectangular or compatible with the selected command. Indoor layout reconstruction assumes a room reasonably approximated by floor, ceiling, and wall planes. Monocular depth may have scale and boundary errors. A Gaussian cloud is a rendering representation and not automatically a metrically complete mesh.')

# Chapter 3
chapter(doc,3,'Background and Core Concepts'); h(doc,'3.1 Equirectangular Projection and Spherical Rays')
para(doc,'An equirectangular panorama maps longitude and latitude to horizontal and vertical pixel coordinates. PanoGS converts pixels to unit ray directions. Given depth d and ray r, the reconstructed point is P = C + d r, where C is the assumed camera center.')
h(doc,'3.2 Monocular Depth'); para(doc,'A monocular depth estimator predicts distance or relative inverse depth from one image. The code tracks whether a result is metric. Depth Anything V2 Metric Indoor is treated as metric; MiDaS is treated as relative depth. Relative output is mapped to a plausible range before back-projection. That mapping is a calibration heuristic, not measured range.')
h(doc,'3.3 3D Gaussian Splatting'); para(doc,'A Gaussian primitive stores a center, anisotropic scale, orientation, opacity, and color features. Its covariance and projected footprint contribute to pixels during alpha compositing. PanoGS initializes primitives from a point cloud and provides CPU and PyTorch rasterizers. The project stores zeroth-order spherical-harmonic color (SH0); it does not implement a full higher-order view-dependent color basis in its main model.')
h(doc,'3.4 LDI and Occlusion Layers'); para(doc,'A Layered Depth Image separates nearer content from a deeper background layer. PanoGS can split a panorama depth map into intervals, inpaint color and depth behind nearer layers, and back-project each active layer. This can reduce visible holes during small movements, but hidden content remains inferred from nearby pixels and priors.')
h(doc,'3.5 Level of Detail'); para(doc,'The custom viewer groups splats in an octree, rejects regions outside the camera frustum, and draws a lower fraction of splats in distant leaves. LOD changes rendering cost, not the underlying geometry.')
table(doc,['Concept','Purpose in PanoGS'],[('Cubemap depth','Run a perspective depth estimator on six overlapping views.'),('Metric depth','Predictions treated as distances in meters, subject to model error.'),('Relative depth','Uncalibrated ordering/inverse depth mapped to an approximate range.'),('LDI','Separate nearer and farther layers and infer some occluded pixels.'),('3D Gaussian','Anisotropic primitive used for splat rendering.'),('LOD','Cull or reduce splats to lower viewer work at distance.')],[1.55,4.9])

# Chapter 4
chapter(doc,4,'Requirements Analysis'); h(doc,'4.1 Functional Requirements')
table(doc,['ID','Requirement','Implementation evidence'],[('F1','Inspect image metadata.','Image I/O and inspect command.'),('F2','Estimate and save depth arrays/visualization.','Depth interface and output writer.'),('F3','Reconstruct panorama geometry.','Cubemap, ray, layout, point-cloud modules.'),('F4','Export Gaussian formats.','Initialization and PLY/SPLAT/HDR SPLAT writers.'),('F5','Open local viewers.','Custom server and bundled SuperSplat.'),('F6','Support video and point-cloud processing.','Video reconstruction and processing modules.'),('F7','Optimize existing Gaussian scenes.','PyTorch training command.'),('F8','Optionally request hosted generation.','World Labs client; needs user key.')],[.5,2.55,3.4],8.7)
h(doc,'4.2 Non-Functional Requirements')
for x in ['Local workflows should not require a paid generation API.','Large panoramas and point clouds require memory and splat-count controls.','Output formats should interoperate with compatible Gaussian viewers.','Depth and geometry uncertainty should be communicated.','Optional viewers and model backends should state browser, hardware, and package needs.']: bullet(doc,x)
h(doc,'4.3 Hardware and Software Requirements')
table(doc,['Area','Evidence / practical requirement'],[('Python package','Python 3.9+; pyproject declares NumPy, Pillow, PyYAML.'),('Computer vision','OpenCV and SciPy used by geometry, inpainting, filtering, video; not declared in pyproject.'),('Neural inference','PyTorch and Transformers used for depth inference and optimization; first run may download weights.'),('Custom viewer','Three.js/WebGL loaded from CDN; bundled SuperSplat is documented as WebGPU-dependent.'),('Hosted generation','Internet, World Labs API key, provider credits.'),('Device','CPU is the local default; GPU is useful for optional deep-learning and external notebook paths.')],[1.45,5.0],8.8)
h(doc,'4.4 Feasibility'); para(doc,'The local path is feasible as a CPU-first reconstruction prototype with optional model downloads and browser viewing. Dependency declarations do not fully describe runtime imports, so packaging should list feature-specific requirements. Accuracy and performance claims require controlled measurements not supplied by the current benchmark alone.')

# Chapter 5
chapter(doc,5,'System Design and Architecture'); h(doc,'5.1 Architecture Overview')
para(doc,'The CLI coordinates image loading, estimator selection, reconstruction, Gaussian initialization, and export. Local Studio accepts uploads and launches reconstruction. The PanoGS viewer serves and renders local splat data; SuperSplat is a separate bundled editor. The World Labs client is an independent remote path.')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(arch),width=Inches(6.35))
p=doc.add_paragraph(style='Caption'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Figure 5.1. PanoGS components and data flow')
h(doc,'5.2 Local Data Flow')
for x in ['The CLI loads the requested source image or video.','A factory creates the estimator; pretrained weights load lazily when inference begins.','Panorama reconstruction samples six overlapping perspective views, estimates and aligns depth, then fuses predictions on spherical rays.','The layout stage filters depth, estimates normals, and can infer background surfaces.','The point cloud is converted into Gaussian parameters and exported.','The local viewer or SuperSplat loads the output scene.']: numbered(doc,x)
h(doc,'5.3 Separate Hosted Path'); para(doc,'The `ai-world` command prepares an LDR preview, uploads it through the World Labs media-asset workflow, requests Marble generation, polls the operation, exports a Gaussian PLY, and writes a manifest. HDR input is copied locally; only the derived display preview is uploaded. This flow consumes provider credits and is separate from local depth inference.')
h(doc,'5.4 Main Modules')
table(doc,['Module','Responsibility'],[('panogs/apps/cli.py','Commands and orchestration.'),('panogs/reconstruction/depth','Depth interface, model factory, cubemap fusion.'),('panogs/reconstruction/panorama','Spherical and layout-guided reconstruction.'),('panogs/reconstruction/pointcloud.py','Depth conversion and back-projection.'),('inpainting.py and ldi.py','Foreground masks, RGB/depth completion, layers.'),('core/gaussian and io/gaussian_ply.py','Gaussian model, initialization, interchange, panoramic repair.'),('training/','Per-scene Gaussian optimization and densification.'),('rendering/ and apps/viewer/','CPU/PyTorch rendering and browser viewing.')],[2.5,3.95],8.7)

# Chapter 6
chapter(doc,6,'Reconstruction Methodology'); h(doc,'6.1 Panorama Input and Depth Estimation')
para(doc,'The documented panorama command defaults to `cubemap_depth_anything`. The factory selects Depth Anything V2 Metric Indoor Small for this name. The cubemap estimator samples six 100-degree perspective faces, yielding about 10 degrees of overlap. Each face is sent to the underlying model. The default face size is 512 by 512 pixels.')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(flow),width=Inches(6.35))
p=doc.add_paragraph(style='Caption'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Figure 6.1. Local panorama reconstruction flow')
h(doc,'6.2 Aligning and Fusing Depth'); para(doc,'For relative-depth faces, PanoGS fits robust affine mappings over shared rays, propagates transforms through the overlap graph, and percentile-normalizes disconnected faces as a fallback. For metric faces, it converts perspective Z depth to radial ray distance, estimates per-face scale corrections in log space, bounds corrections, and anchors horizontal faces. It then reprojects predictions to panorama rays and feathers overlap zones.')
h(doc,'6.3 Room Priors and Depth Cleanup'); para(doc,'The layout path applies morphology-based depth regularization, RGB-guided joint bilateral smoothing, and edge sharpening. It intersects rays with an assumed Manhattan room envelope and computes planar normals. Foreground masks come from depth differences against a local background estimate. These are image-processing and geometry rules, not learned semantic object segmentation.')
h(doc,'6.4 Back-Projection and Gaussian Initialization'); para(doc,'PanoGS generates spherical unit rays and computes P = C + d r. It filters invalid or out-of-range depths and can remove large relative discontinuities. Gaussian scale derives from local nearest-neighbor spacing with outlier caps and panorama angular-footprint bounds. If normals exist, rotations and anisotropic scales align to tangent frames. Opacity is in logit form; color is mapped to SH0.')
h(doc,'6.5 Filling Black or Empty Regions'); para(doc,'Several deterministic completion methods are used. Telea or Navier-Stokes inpainting propagates nearby colors into masks. A nearest-valid-pixel fallback replaces residual black pixels inside an inpaint mask. Background depth is propagated from valid pixels and smoothed in inverse-depth space; a room envelope can constrain it. LDI repeats RGB, radiance, and depth inpainting between layers. Panoramic Gaussian repair can fill missing spherical grid cells from neighboring samples and add a simple nadir floor cap.')
para(doc,'These methods produce plausible fills from nearby evidence and geometric priors; they cannot recover actual hidden appearance or geometry. Unmasked black pixels in the input remain image content. A dark empty viewer background is a rendering backdrop, not a generated room surface.')
h(doc,'6.6 Video and Novel-View Prototypes'); para(doc,'The video path selects sharp frames using Laplacian variance, estimates depth, matches ORB features, fits an essential matrix with RANSAC, estimates relative pose, back-projects frame points, then merges them using voxel filtering and PCA normals. It uses approximate intrinsics and fallback motion values, so trajectory scale and drift require validation. The separate virtual-camera synthesis module renders point-cloud buffers and fills internal holes using classical inpainting; “generative” in its comments does not mean a neural image-generation model is called.')

# Chapter 7
chapter(doc,7,'Implementation and Technology Stack'); h(doc,'7.1 AI Models and Weight Loading')
table(doc,['Model','Use','Loading / status'],[('Depth Anything V2 Metric Indoor Small/Base/Large','Monocular metric depth; Small is panorama default.','Hugging Face Transformers depth pipeline; weights load lazily.'),('MiDaS Small, DPT Hybrid, DPT Large','Alternative relative inverse depth.','PyTorch Hub loads intel-isl/MiDaS and transforms.'),('Synthetic room / gradient','Deterministic test/demo depth.','NumPy formulas; not learned.'),('Marble 1.x','Optional hosted image-to-world generation.','World Labs API; no local checkpoint.'),('InfiniSplat RGB','External single-crop notebook baseline.','Separate Kaggle GPU environment, not core CLI.'),('DA360 via BLUNT','External panorama comparison notebook.','Separate repository and checkpoint workflow.')],[1.55,2.1,2.8],8.1)
para(doc,'Transformers is a software library used to load and run the Depth Anything vision model; it is not itself the model. For each image the estimator preprocesses, performs a forward pass, retrieves predicted depth, resizes when needed, and returns a float32 DepthResult. PanoGS does not update these pretrained depth weights during image reconstruction.')
h(doc,'7.2 Training and Optimization'); para(doc,'The `train` command loads an initial Gaussian PLY and samples perspective crops from the same image or panorama. A custom differentiable PyTorch rasterizer renders one view at a time. The loss combines mean absolute RGB error with 1 minus SSIM. Adam updates position, log scale, quaternion rotation, opacity logits, and SH0 color. Gradient thresholds trigger cloning or splitting; low-opacity or oversized splats can be pruned, and opacity can be reset.')
para(doc,'This is per-scene optimization, not depth-model training or training a general-purpose image-to-3D model. The sampled views share the panorama capture center and provide no independent parallax. Defaults (150 iterations and 128-pixel views) are development-scale and do not establish convergence.')
h(doc,'7.3 API Keys and External Services'); para(doc,'Only the World Labs client reads an application key: `WLT_API_KEY`, with `WORLDLABS_API_KEY` as an alternate environment variable. It sends the value in the `WLT-Api-Key` request header. No key is hard-coded in the local pipeline. Model weights are obtained through their model ecosystems; notebooks download external checkpoints. Hosted generation needs network access and provider credits.')
h(doc,'7.4 Libraries and Their Uses')
table(doc,['Library / tool','Use'],[('NumPy','Arrays, vectorized geometry, rays, color/depth operations.'),('Pillow','Image loading, conversion, preview creation.'),('PyYAML','Configuration parsing.'),('PyTorch','MiDaS inference and Gaussian scene optimization.'),('Transformers','Hugging Face depth-estimation pipeline for Depth Anything V2.'),('OpenCV','Filtering, morphology, inpainting, video, ORB and RANSAC.'),('SciPy','KD trees, distance transforms, Gaussian filters, interpolation.'),('Three.js / OrbitControls','Custom WebGL viewer, loaded from CDN.'),('SuperSplat','Bundled PlayCanvas editor for Gaussian PLY assets.'),('pytest','Test framework; source tests are present.')],[1.7,4.75],8.5)
para(doc,'The declared package dependencies are narrower than the imports. OpenCV, SciPy, PyTorch, and Transformers need to be installed for features using them. Packaging should declare feature groups or document these requirements.')
h(doc,'7.5 Export and Viewing'); para(doc,'Gaussian PLY stores centers, log scales, rotations, opacity logits, and SH0 color features. SuperSplat-compatible export adjusts coordinate and color conventions. `.splat` is a viewer format; `.hdrsplat` retains linear radiance for HDR controls. PanoGS also has a Three.js/WebGL viewer with frustum culling and octree LOD. SuperSplat is a separate editor and may require WebGPU.')

# Chapter 8
chapter(doc,8,'Testing and Evaluation'); h(doc,'8.1 Repository Tests')
para(doc,'The repository contains 18 test modules and 90 test functions covering CLI, configuration, depth factories, image loading, Gaussian math, PLY I/O, point-cloud processing, inpainting, LDI, synthesis, rendering, training utilities, video, compression, benchmarks, rays, and the viewer. These counts show intended source coverage; this report does not claim the tests were executed or passed.')
h(doc,'8.2 What Current Metrics Measure'); para(doc,'The training loop logs L1, SSIM, and PSNR for the current sampled view. These measure image agreement on panorama crops used during optimization. Since those views are generated from the input and share the capture center, they are not independent held-out tests of depth or geometry.')
para(doc,'The benchmark reports file size, splat count, load time, and CPU render timing on an orbit path. Its PSNR and approximate SSIM compare adjacent rendered frames, not ground-truth images. They are frame-to-frame similarity indicators, not reconstruction accuracy.')
table(doc,['Measure','Current status','Needed for defensible claim'],[('Depth accuracy','No GT depth evaluation found.','Ground truth; AbsRel, RMSE, threshold accuracy; align relative scale.'),('3D geometry','No reference scan/mesh comparison.','Completeness, point-to-surface or Chamfer distance, normal error.'),('Novel views','Training uses source-derived crops; benchmark compares adjacent renders.','Real held-out camera images; fixed protocol; PSNR, SSIM, visual review.'),('Runtime','CPU load/render timing exists.','Same device, resolution, path, warm-up, memory, repeated runs.'),('Model selection','No controlled winner study found.','Same scenes/settings; quality, speed, failures, ablation study.')],[1.3,2.35,2.8],8.2)
h(doc,'8.3 How to Determine Which Method Is Best'); para(doc,'Define “best” for the intended use: metric distance, visual reproduction, browser speed, size, or quality after translation. Compare candidates on the same scenes, preprocessing, resolution, hardware, and fixed held-out views. Change one component at a time and include difficult cases such as thin objects, blank walls, poles, and occlusions. No model should be declared best from subjective preference or training loss alone.')

# Chapter 9
chapter(doc,9,'Results and Discussion'); h(doc,'9.1 Implemented Capabilities')
para(doc,'The source tree includes local panorama reconstruction, image inspection and depth export, point-cloud reconstruction and processing, Gaussian initialization and export, CPU/PyTorch renderers, compression, scene training, video reconstruction, local viewing, a SuperSplat editor bundle, and an optional Marble API client. These are source-level capabilities; this report does not claim runtime verification of every command.')
h(doc,'9.2 Documented Example Outputs'); para(doc,'The single-panorama guide records example Brown Photostudio outputs of about 2.40 million Gaussians at original resolution and about 0.61 million for a 1024-pixel preview. It notes 512-pixel cubemap face inference and artifacts when translating away from the capture center. The report does not independently reproduce those outputs; the figures are attributed to project documentation, not presented as fresh measurements.')
h(doc,'9.3 Interpretation'); para(doc,'PanoGS integrates pretrained monocular depth with panorama-aware geometry, room-layout priors, Gaussian export, and local viewing. The result is best described as an inferred scene for interactive look-around. The single-panorama route does not add observations behind occluders or at new camera positions. Marble output is an AI-generated interpretation and should be distinguished from local reconstruction.')

# Chapter 10
chapter(doc,10,'Limitations, Future Scope and Conclusion'); h(doc,'10.1 Current Limitations')
for x in ['One panorama supplies no parallax; hidden geometry and off-center movement are uncertain.','Depth can be wrong at boundaries, poles, reflections, or textureless surfaces.','Inpainting and room priors produce plausible fills, not verified content.','Video motion uses approximate intrinsics and fallback translations.','The benchmark does not measure ground-truth reconstruction accuracy.','Training optimizes scene appearance from source-derived views at small default resolution.','The package dependency list omits several runtime libraries.','Hosted Marble depends on network, a key, and provider credits.']: bullet(doc,x)
h(doc,'10.2 Future Work')
for x in ['Evaluate depth and geometry against public datasets and reference scans.','Build a held-out multi-view comparison with InfiniSplat and other baselines.','Improve video poses with calibrated intrinsics, robust full-pose estimation, and loop closure.','Make model-download failures and fallback behavior explicit and test actual fallback inference.','Declare optional dependency groups and verify clean installs.','Profile memory, sorting, and viewer LOD on fixed scenes.','Label inferred completion separately from measured geometry.','Add viewpoint-aware quality metrics and evidence capture to releases.']: numbered(doc,x)
h(doc,'10.3 Conclusion'); para(doc,'PanoGS is a broad prototype for converting one panorama into an interactive Gaussian scene. The central pipeline estimates depth on overlapping perspective crops, fuses it with spherical geometry and indoor priors, back-projects pixels, initializes Gaussian primitives, and exports viewer-compatible files. Classical filtering and inpainting can reduce seams and holes; local viewers and LOD help inspect large outputs.')
para(doc,'The project does not retrain its pretrained depth models at image-submission time, and current code does not establish that its reconstruction is more accurate than other systems. Ground-truth evaluation, held-out views, documented hardware, and repeatable comparisons are needed for that claim.')

# References
doc.add_page_break(); p=doc.add_paragraph(style='Heading 1'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('REFERENCES')
refs=['[1] PanoGS project source and documentation: README.md; docs/single_panorama_pipeline.md; docs/depth.md; docs/optimization.md; docs/performance.md; panogs/ source modules.','[2] B. Kerbl et al., “3D Gaussian Splatting for Real-Time Radiance Field Rendering,” ACM Transactions on Graphics, 2023. https://doi.org/10.1145/3592433','[3] L. Yang et al., “Depth Anything V2,” arXiv:2406.09414, 2024. https://arxiv.org/abs/2406.09414','[4] R. Ranftl, A. Bochkovskiy, and V. Koltun, “Vision Transformers for Dense Prediction,” ICCV 2021. https://arxiv.org/abs/2103.13413','[5] “PanoDreamer: 3D Panorama Synthesis from a Single Image,” arXiv:2412.04827. https://arxiv.org/abs/2412.04827','[6] Hugging Face Transformers pipeline documentation. https://huggingface.co/docs/transformers/main_classes/pipelines','[7] OpenCV documentation for inpainting, morphology, and video APIs. https://docs.opencv.org/','[8] PlayCanvas SuperSplat project and license material bundled in panogs/apps/viewer.','[9] World Labs Marble API integration as implemented in panogs/generation/worldlabs.py.','[10] zju3dv/InfiniSplat, used only by a separate baseline notebook. https://github.com/zju3dv/InfiniSplat','[11] SonnyC56/BLUNT and DA360 workflow, used only by a separate Kaggle notebook. https://github.com/SonnyC56/blunt']
for x in refs:
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.28); p.paragraph_format.first_line_indent=Inches(-.28); p.paragraph_format.space_after=Pt(7); p.add_run(x)

# Appendix A
doc.add_page_break(); p=doc.add_paragraph(style='Heading 1'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('APPENDIX A  RUNNING THE PROJECT')
h(doc,'A.1 Common Local Commands')
for x in ['pip install -e ".[dev]" (then install optional runtime dependencies)','panogs inspect path/to/panorama.jpg','panogs depth path/to/image.jpg --model depth_anything_v2','panogs reconstruct path/to/panorama.hdr --model cubemap_depth_anything --shape hybrid -o output/scene.ply','panogs video path/to/walkthrough.mp4 -o output/video_scene.ply','panogs train path/to/panorama.jpg --init-ply output/scene_gaussians.ply -o output/optimized.ply','panogs view output/scene.splat','panogs benchmark output/scene_gaussians.ply']:
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.3); p.paragraph_format.first_line_indent=Inches(0); p.paragraph_format.space_after=Pt(5); r=p.add_run(x); r.font.name='Consolas'; r.font.size=Pt(9)
h(doc,'A.2 World Labs Key'); p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.3); p.paragraph_format.first_line_indent=Inches(0); r=p.add_run('$env:WLT_API_KEY = "your-api-key"'); r.font.name='Consolas'; r.font.size=Pt(9)
para(doc,'Then run `panogs ai-world path/to/panorama.hdr --open-viewer`. Do not commit a real key to source control.')
h(doc,'A.3 Operational Notes')
for x in ['Depth weights load lazily and may download on first inference.','Large panoramas and multi-million-splat scenes can require substantial time and memory.','Use preview resolution while iterating; preserve the original HDR if radiance matters.','If default depth inference fails, the CLI does not guarantee an automatic MiDaS retry; choose a model and install its dependencies explicitly.']: bullet(doc,x)

# Appendix B
doc.add_page_break(); p=doc.add_paragraph(style='Heading 1'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('APPENDIX B  REPOSITORY TEST COVERAGE')
para(doc,'The table lists test modules and source-level function counts. Counts are not pass/fail results.')
mods=[('test_benchmark.py',2),('test_cli.py',12),('test_compression.py',3),('test_config.py',3),('test_depth.py',7),('test_gaussian.py',11),('test_image_loader.py',8),('test_inpainting.py',3),('test_ldi.py',2),('test_ply.py',3),('test_pointcloud.py',5),('test_processing.py',5),('test_rays.py',3),('test_renderer.py',5),('test_synthesis.py',5),('test_training.py',5),('test_video.py',3),('test_viewer.py',5)]
table(doc,['Test module','Functions'],mods+[('Total',90)],[4.8,1.65],9)
para(doc,'Before submission, run the suite in the intended environment, record the command and result, and attach genuine application screenshots or generated artifacts. A passing unit suite is not evidence of geometric accuracy; that requires ground-truth evaluation.')

doc.core_properties.title='PanoGS An Engine for Converting Single Panoramic Images into 3D Gaussian Scenes'
doc.core_properties.subject='Academic project report on panorama-to-3D Gaussian reconstruction'
doc.core_properties.author='Abel Bobby; Siddharth Gupta; Aditya Sharma'
doc.core_properties.keywords='PanoGS, panorama, monocular depth, 3D Gaussian Splatting'
doc.save(DOCX)
print(DOCX)
