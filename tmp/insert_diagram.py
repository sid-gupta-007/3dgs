from pathlib import Path
from docx import Document
from docx.shared import Inches
root=Path(r'C:\Users\Siddharth Gupta\Desktop\CODES-ALL\3dg')
docx=root/'output'/'PanoGS_Project_Report.docx'; image=root/'output'/'PanoGS_Architecture_Diagram.png'
d=Document(docx)
cell=next(c for t in d.tables for row in t.rows for c in row.cells if 'PASTE PROJECT IMAGE HERE' in c.text)
cell.text=''
p=cell.paragraphs[0]; p.alignment=1
p.add_run().add_picture(str(image),width=Inches(6.15))
d.save(docx); print(docx)
