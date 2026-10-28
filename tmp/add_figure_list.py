from pathlib import Path
from docx import Document
from docx.shared import Pt,Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
root=Path(r'C:\Users\Siddharth Gupta\Desktop\CODES-ALL\3dg')
path=root/'output'/'PanoGS_Project_Report.docx'; d=Document(path)
break_before_ch1=next(p for p in d.paragraphs if p.style.name=='Heading 1' and p.text.startswith('CHAPTER 1\n'))._p.getprevious()
items=['Figure 5.2  Project interface or architecture screenshot','Figure 6.2  Input panorama, depth estimate, and reconstruction','Figure 9.1  PanoGS scene result and evaluation evidence']
for text in items:
    p=d.add_paragraph(); p.paragraph_format.left_indent=Inches(.22); p.paragraph_format.first_line_indent=Inches(0); p.paragraph_format.space_after=Pt(5); p.add_run(text)
    e=p._p; e.getparent().remove(e); break_before_ch1.addprevious(e)
d.save(path)
print(path)
